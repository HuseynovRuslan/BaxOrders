"""Build the order workbook the objects fill in, with the catalogue built into it as one dropdown.

Why this exists. Every name in the workbook used to be typed by hand, and the reconciliation of October 2026
showed what that costs: one article written «Xlor Toz 90% qranul» on 25 rows, «Xlor toz qranul 90% Pooltech»
on 9 and «Xlor toz qranul 90%» on 7, and the same for nine other articles; the same drift in sizes
(«115» against «115x10x22,23») and in model codes («Emaf A300» against «EMAF A 300»). Each spelling then has
to be recognised by hand at import time, and a wrong guess puts one article in the catalogue twice, after
which no report can add it up.

So nothing about the article is typed here. One dropdown offers the catalogue's variants - «Laqonda daşı
almaz · 115 MM · Staunch · Ədəd» - and choosing one fills in the product, the size, its unit, the model code
and the order unit together. Those five columns are formulas on a hidden, protected range: they are what the
import reads, and they cannot disagree with each other or with the catalogue because nobody writes them.

What is left to fill in is the order itself: how many, at what price, for which object, when, from whom, by
whom, and whether it is the monthly order or on top of it. The first two are checked as numbers, the date as
a date, and the rest come from their own dropdowns.

What it cannot do. The lists are a snapshot of the catalogue on the day the file is built, so a variant that
does not exist yet cannot be ordered through this template. That is the point: a new article, or a new size
of an old one, is a catalogue decision. Add it to the catalogue and rebuild the file - it takes one command.

Usage:
    python build_template.py <reference-csv-dir> <output.xlsx>
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SHEET = "Aylıq"          # ImportSheets.Authoritative - the import reads this sheet and no other
REFERENCE = "Kataloq"
GUIDE = "Təlimat"
HEADER_ROW = 3
FIRST_ROW = HEADER_ROW + 1
LAST_ROW = 1003
CHOICE = "Mal seçimi"    # not an ImportHeaders name, so the import ignores this column

# The columns, in order. "derived" ones are formulas filled from the chosen variant and hidden: they are the
# columns the import actually reads, and showing them again would only repeat what the choice already says.
# «Ana/Orta/Alt kateqoriya» are absent on purpose - they are not required and the catalogue knows a product's
# category without being told.
COLUMNS = [
    ("№", 7, "auto"),
    (CHOICE, 62, "variant"),
    ("Malın adı", 36, "derived"),
    ("Ölçü", 12, "derived"),
    ("Vahid", 14, "derived"),
    ("Model/Kod", 18, "derived"),
    ("Miqdarı", 11, "quantity"),
    ("Ö/v", 12, "derived"),
    ("Qiyməti", 12, "money"),
    ("Cəmi", 13, "auto"),
    ("Şirkət", 28, "supplier"),
    ("Obyekt", 26, "location"),
    ("Tarix", 12, "date"),
    ("Cavabdeh şəxs", 22, "responsible"),
    ("Növ", 16, "linetype"),
]
# Which Kataloq column each derived column reads, 1-based within the variant table A:F.
DERIVED_FROM = {"Malın adı": 2, "Ölçü": 3, "Vahid": 4, "Model/Kod": 5, "Ö/v": 6}
LINE_TYPES = ["Aylıq sifariş", "Aylıqdan kənar"]

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
PICK_FILL = PatternFill("solid", fgColor="2E5C9A")
AUTO_FILL = PatternFill("solid", fgColor="EFEFEF")
THIN = Side(style="thin", color="BFBFBF")


def read_column(path, index=0):
    with open(path, encoding="utf-8-sig") as fh:
        rows = list(csv.reader(fh))
    return [r[index] for r in rows[1:] if r and r[index].strip()]


def read_variants(path):
    """One row per orderable variant: what a person reads, and the five values the import needs."""
    out = []
    with open(path, encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            product = r["product"].strip()
            if not product:
                continue
            spec = r["spec"].strip()
            # A specification carries either a unit or a descriptor, never both, and the import reads the
            # descriptor out of the same «Vahid» column - so one column holds whichever this variant has.
            unit_short = r["spec_unit_code"].strip() or r["descriptor"].strip()
            unit_value = r["spec_unit_name"].strip() or r["descriptor"].strip()
            model = r["model"].strip()
            order_unit = r["order_unit"].strip()

            label = [product]
            if spec:
                label.append(f"{spec} {unit_short}".strip())
            if model:
                label.append(model)
            if order_unit:
                label.append(order_unit)
            out.append({"label": " · ".join(label), "product": product, "spec": spec,
                        "unit": unit_value, "model": model, "order_unit": order_unit})

    seen = {}
    for v in out:
        if v["label"] in seen:
            raise SystemExit(f"iki variant eyni adla görünür: «{v['label']}» — siyahı birmənalı olmalıdır")
        seen[v["label"]] = v
    return out


def build(reference_dir, out_path):
    ref = Path(reference_dir)
    variants = read_variants(ref / "ref-variant.csv")
    locations = read_column(ref / "ref-loc.csv")
    suppliers = read_column(ref / "ref-sup.csv")
    responsible = read_column(ref / "ref-resp.csv")

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    cat = wb.create_sheet(REFERENCE)
    guide = wb.create_sheet(GUIDE)

    # ---------------------------------------------------------------- the lists the dropdowns read
    for column, title in enumerate(["Seçim", "Malın adı", "Ölçü", "Vahid", "Model/Kod", "Ö/v"], start=1):
        cat.cell(1, column, title).font = Font(bold=True)
    for i, v in enumerate(variants, start=2):
        for column, key in enumerate(["label", "product", "spec", "unit", "model", "order_unit"], start=1):
            cat.cell(i, column, v[key])

    lists = {"Seçim": f"={REFERENCE}!$A$2:$A${len(variants) + 1}"}
    for column, title, values in (
        (8, "Obyekt", locations), (10, "Şirkət", suppliers), (12, "Cavabdeh şəxs", responsible),
        (14, "Növ", LINE_TYPES),
    ):
        cat.cell(1, column, title).font = Font(bold=True)
        for i, value in enumerate(values, start=2):
            cat.cell(i, column, value)
        letter = get_column_letter(column)
        lists[title] = f"={REFERENCE}!${letter}$2:${letter}${len(values) + 1}"
    for letter, width in (("A", 64), ("B", 36), ("C", 10), ("D", 14), ("E", 18), ("F", 12),
                          ("H", 26), ("J", 28), ("L", 22), ("N", 16)):
        cat.column_dimensions[letter].width = width
    cat.freeze_panes = "A2"

    # ---------------------------------------------------------------- the entry sheet
    ws["A1"] = "BAX Orders — aylıq sifariş cədvəli"
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    ws["A2"] = ("«Mal seçimi» sütununda malı siyahıdan seçin — ölçü, vahid və model özü dolur. "
                "Boz sütunlar hesablanır. Əldən yazılan ad qəbul edilmir.")
    ws["A2"].font = Font(italic=True, size=9, color="7F7F7F")

    at = {}
    for index, (title, width, kind) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        at[title] = letter
        cell = ws.cell(HEADER_ROW, index, title)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PICK_FILL if kind == "variant" else HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        ws.column_dimensions[letter].width = width
        if kind == "derived":
            # Hidden, not deleted: the import reads these columns, and the chosen line already shows a
            # person everything they say.
            ws.column_dimensions[letter].hidden = True
    ws.row_dimensions[HEADER_ROW].height = 30
    ws.freeze_panes = f"{at[CHOICE]}{FIRST_ROW}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(COLUMNS))}{LAST_ROW}"

    pick = at[CHOICE]
    for row in range(FIRST_ROW, LAST_ROW + 1):
        ws[f'{at["№"]}{row}'] = f'=IF(${pick}{row}="","",ROW()-{HEADER_ROW})'
        for title, lookup in DERIVED_FROM.items():
            # An empty cell in the lookup range reads back as 0, so the zero becomes a blank again; no real
            # size, unit or model code can ever be 0, which makes the test safe.
            found = f'VLOOKUP(${pick}{row},{REFERENCE}!$A:$F,{lookup},FALSE)'
            ws[f"{at[title]}{row}"] = f'=IFERROR(IF({found}=0,"",{found}),"")'
        ws[f'{at["Cəmi"]}{row}'] = (
            f'=IF(AND(${at["Miqdarı"]}{row}<>"",${at["Qiyməti"]}{row}<>""),'
            f'${at["Miqdarı"]}{row}*${at["Qiyməti"]}{row},"")')

    validations = {
        "variant": DataValidation(type="list", formula1=lists["Seçim"], allow_blank=True,
                                  errorTitle="Bu mal kataloqda yoxdur",
                                  error="Malı siyahıdan seçin. Lazım olan mal və ya ölçü siyahıda yoxdursa, "
                                        "əvvəlcə kataloqa əlavə edilməlidir — administratora bildirin."),
        "location": DataValidation(type="list", formula1=lists["Obyekt"], allow_blank=True,
                                   errorTitle="Obyekt tanınmır", error="Obyekti siyahıdan seçin."),
        "supplier": DataValidation(type="list", formula1=lists["Şirkət"], allow_blank=True,
                                   errorTitle="Şirkət tanınmır",
                                   error="Təchizatçını siyahıdan seçin, ya da xananı boş buraxın."),
        "responsible": DataValidation(type="list", formula1=lists["Cavabdeh şəxs"], allow_blank=True,
                                      errorTitle="Cavabdeh şəxs tanınmır",
                                      error="Cavabdeh şəxsi siyahıdan seçin."),
        "linetype": DataValidation(type="list", formula1=lists["Növ"], allow_blank=True,
                                   errorTitle="Növ tanınmır",
                                   error="«Aylıq sifariş» ya da «Aylıqdan kənar» seçin."),
        "quantity": DataValidation(type="decimal", operator="greaterThan", formula1=0, allow_blank=True,
                                   errorTitle="Miqdar düzgün deyil",
                                   error="Miqdar sıfırdan böyük rəqəm olmalıdır."),
        "money": DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0, allow_blank=True,
                                errorTitle="Qiymət düzgün deyil",
                                error="Qiymət mənfi ola bilməz. Qiymət yoxdursa xananı boş buraxın."),
        "date": DataValidation(type="date", operator="greaterThanOrEqual", formula1="DATE(2020,1,1)",
                               allow_blank=True, errorTitle="Tarix düzgün deyil",
                               error="Tarixi gün/ay/il formatında yazın."),
    }
    for dv in validations.values():
        dv.showErrorMessage = True
        dv.errorStyle = "stop"       # refuse the value, never merely warn: that is the whole point
        ws.add_data_validation(dv)

    for index, (title, _, kind) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        if kind in validations:
            validations[kind].add(f"{letter}{FIRST_ROW}:{letter}{LAST_ROW}")
        for row in range(FIRST_ROW, LAST_ROW + 1):
            cell = ws.cell(row, index)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if kind in ("auto", "derived"):
                cell.fill = AUTO_FILL
                cell.protection = Protection(locked=True)
            else:
                cell.protection = Protection(locked=False)
            if kind == "date":
                cell.number_format = "DD.MM.YYYY"
            elif title in ("Qiyməti", "Cəmi"):
                cell.number_format = "#,##0.00"
            elif kind == "quantity":
                cell.number_format = "#,##0.###"
            elif title == "№":
                cell.number_format = "0"

    # Protection keeps the formulas safe; everything a person fills stays unlocked. No password, so whoever
    # needs to restructure the sheet can simply turn it off.
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False

    # ---------------------------------------------------------------- the instructions
    guide.column_dimensions["A"].width = 112
    text = [
        ("BAX Orders — sifariş cədvəlinin doldurulması və yüklənməsi", "title"),
        ("", None),
        ("Bu fayl nəyə görə belədir", "head"),
        ("Əvvəl adlar, ölçülər və modellər əldən yazılırdı. Nəticədə bir mal bir vərəqdə bir neçə cür "
         "görünürdü — «Xlor Toz 90% qranul» 25 sətirdə, «Xlor toz qranul 90% Pooltech» 9 sətirdə, "
         "«Xlor toz qranul 90%» 7 sətirdə; ölçü gah «115», gah «115x10x22,23». Hər yazılışı idxal zamanı "
         "əl ilə tanıtmaq lazım gəlir və səhv tanıdılsa bir mal kataloqda iki dəfə yaranır, ondan sonra "
         "heç bir hesabat onu toplaya bilmir.", None),
        ("Bu şablonda malın heç bir xüsusiyyəti yazılmır — biri seçilir, qalanı özü dolur.", None),
        ("", None),
        ("Necə doldurulur", "head"),
        ("1. «Mal seçimi» xanasına klikləyin, sağda çıxan oxu basın və siyahıdan seçin. Siyahıda hər sətir "
         "belədir:  Laqonda daşı almaz · 115 MM · Staunch · Ədəd  —  yəni mal, ölçü, model və ölçü vahidi.", None),
        ("2. Seçdikdən sonra «Malın adı», «Ölçü», «Vahid», «Model/Kod» və «Ö/v» sütunları özü dolur. "
         "Onlar gizlədilib, çünki seçdiyiniz sətir onsuz da hamısını göstərir.", None),
        ("3. «Miqdarı» yazın — sıfırdan böyük rəqəm.", None),
        ("4. «Qiyməti» yazın. Mal daxili anbardan verilibsə xananı BOŞ buraxın.", None),
        ("5. «Şirkət», «Obyekt», «Cavabdeh şəxs», «Növ» — siyahıdan seçilir.", None),
        ("6. «Tarix» — sifarişin aid olduğu AYI göstərir, gün əhəmiyyətsizdir.", None),
        ("7. «№» və «Cəmi» özü hesablanır. Onlara yazmayın.", None),
        ("", None),
        ("Dəyişməməli olanlar", "head"),
        ("Vərəqin adı «Aylıq» olmalıdır — idxal yalnız bu adda vərəqi oxuyur, başqa vərəq oxunmur. "
         "Sütun adlarını da dəyişməyin: idxal sütunu adına görə tanıyır.", None),
        ("", None),
        ("Lazım olan mal və ya ölçü siyahıda yoxdursa", "head"),
        ("Əldən yazmayın — Excel onsuz da qəbul etməyəcək. Yeni mal, yaxud köhnə malın yeni ölçüsü kataloq "
         "qərarıdır: administratora bildirin, kataloqa əlavə edilsin, sonra bu şablon yenidən yaradılsın. "
         "Siyahı kataloqun fayl yaradılan günə olan vəziyyətidir.", None),
        ("", None),
        ("Necə yüklənir", "head"),
        ("1. Saytda «İdxal» bölməsini açın.", None),
        ("2. Faylı seçib yükləyin. Sistem sətirləri oxuyur və tanımadığı dəyərləri sizə göstərir.", None),
        ("3. Tanınmayan dəyər varsa, onu nəzərdən keçirmə ekranında həll edin.", None),
        ("4. Təsdiq ekranında hansı sifarişlərin yaranacağını yoxlayın və təsdiqləyin.", None),
        ("5. Sifarişlər «Qaralama» kimi yaranır; təsdiq axını saytda davam edir.", None),
        ("", None),
        ("Qeyd: eyni fayl iki dəfə yüklənə bilməz — sistem təkrarı tanıyır və rədd edir.", None),
    ]
    row = 1
    for line, style in text:
        cell = guide.cell(row, 1, line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if style == "title":
            cell.font = Font(bold=True, size=14, color="1F3864")
        elif style == "head":
            cell.font = Font(bold=True, size=11, color="1F3864")
        row += 1

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"variants": len(variants), "products": len({v["product"] for v in variants}),
            "locations": len(locations), "suppliers": len(suppliers), "responsible": len(responsible)}


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    counts = build(sys.argv[1], sys.argv[2])
    print(f"{sys.argv[2]} yazıldı — " + ", ".join(f"{k}: {v}" for k, v in counts.items()))
